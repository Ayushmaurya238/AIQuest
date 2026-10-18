"""Build Track 2 coverage, correctness, and reproducibility evidence."""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import platform
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from pipeline.common import ensure_dirs, load_config, write_json


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _csv_count(path: Path) -> int | None:
    if not path.is_file():
        return None
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return max(0, sum(1 for _ in csv.reader(handle)) - 1)


def _workbook_check(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {"available": False, "error": "workbook_missing"}
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        expected_sheets = ["Run Summary", "Provider Rates 1", "Unresolved Rates"]
        sheets = book.sheetnames
        if "Provider Rates 1" not in sheets or "Unresolved Rates" not in sheets:
            return {"available": True, "sheet_names": sheets, "error": "expected_sheet_missing"}
        rates = book["Provider Rates 1"]
        header = next(rates.iter_rows(min_row=1, max_row=1, values_only=True), ())
        required_keys = ["NPI", "TIN", "Address", "Address 2", "City", "State", "ZIP", "Billing Code"]
        missing = [key for key in required_keys if key not in header]
        if missing:
            return {"available": True, "sheet_names": sheets, "missing_key_columns": missing}
        positions = [header.index(key) for key in required_keys]
        seen: set[tuple[Any, ...]] = set()
        duplicate_count = 0
        row_count = 0
        for row in rates.iter_rows(min_row=2, values_only=True):
            row_count += 1
            key = tuple(row[position] or "" for position in positions)
            if key in seen:
                duplicate_count += 1
            else:
                seen.add(key)
        unresolved_rows = sum(1 for _ in book["Unresolved Rates"].iter_rows(min_row=2, values_only=True))
        return {
            "available": True,
            "sheet_names": sheets,
            "expected_sheets_present": sheets == expected_sheets,
            "provider_rate_rows": row_count,
            "unique_provider_rate_keys": len(seen),
            "duplicate_provider_rate_keys": duplicate_count,
            "unresolved_workbook_rows": unresolved_rows,
        }
    finally:
        book.close()


def _ratio(numerator: Any, denominator: Any) -> float | None:
    try:
        n = float(numerator)
        d = float(denominator)
        return round(n / d, 6) if d else None
    except (TypeError, ValueError):
        return None


def run(config: dict[str, Any], logger: Any = None, external_sources: list[Path] | None = None) -> dict[str, Any]:
    dirs = ensure_dirs(config)
    project_root = Path(config.get("_project_root", Path.cwd())).resolve()
    processed = dirs["processed"]
    output = dirs["output"]
    sources = config.get("sources", {})
    pipeline_cfg = config.get("pipeline", {})

    c1 = _json(processed / "c1_manifest.json")
    c2 = _json(processed / "c2_manifest.json")
    c4 = pd.read_csv(processed / "c4_file_status.csv", dtype=str, keep_default_na=False) if (processed / "c4_file_status.csv").is_file() else pd.DataFrame()
    c5 = _json(processed / "c5_manifest.json")
    c6 = _json(processed / "c6_manifest.json")
    cms = _json(processed / "cms_source_audit.json")
    quality = pd.read_csv(processed / "file_quality_metrics.csv", dtype=str, keep_default_na=False) if (processed / "file_quality_metrics.csv").is_file() else pd.DataFrame()
    workbook = _workbook_check(output / "provider_rates.xlsx")

    c4_rows = int(pd.to_numeric(c4.get("rate_rows_extracted", pd.Series(dtype=str)), errors="coerce").fillna(0).sum()) if not c4.empty else 0
    c4_complete = int((c4.get("status", pd.Series(dtype=str)) == "complete").sum()) if not c4.empty else 0
    selected_files = int(c2.get("files_selected_for_extraction", 0) or 0)
    selected_plans = int(c2.get("unique_explicit_ppo_plans", 0) or 0)
    total_pairs = int(c5.get("rate_npi_tin_pairs", 0) or 0)
    candidate_pairs = int(c5.get("pairs_with_address_candidates", 0) or 0)
    verified_pairs = int(c5.get("verified_pair_address_rows", 0) or 0)
    unresolved = int(c6.get("unresolved_or_unaddressed_numeric_rate_rows", 0) or 0)
    provider_rate_rows = int(c6.get("provider_rate_rows", 0) or 0)
    mapped_rows = int(c6.get("billing_code_specialty_matches", 0) or 0)
    unmapped_rows = int(c6.get("unmapped_specialty_rate_rows", 0) or 0)
    quality_score = quality.iloc[0].get("quality_score", "") if not quality.empty else ""
    quality_grade = quality.iloc[0].get("grade", "") if not quality.empty else ""

    checks: dict[str, Any] = {
        "C1_index_manifest_complete": c1.get("status") == "complete",
        "C2_plan_classification_complete": c2.get("status") == "complete",
        "C4_has_complete_rate_file": c4_complete > 0,
        "workbook_rate_row_count_matches_C6_manifest": workbook.get("provider_rate_rows") == provider_rate_rows,
        "provider_rate_keys_are_unique": workbook.get("duplicate_provider_rate_keys") == 0,
        "unresolved_sheet_matches_C6_manifest": workbook.get("unresolved_workbook_rows") == unresolved,
        "expected_workbook_sheets_present": workbook.get("expected_sheets_present") is True,
        "C5_address_mapping_completed": c5.get("status") == "complete",
        "C6_workbook_completed": c6.get("status") == "complete" and workbook.get("available") is True,
    }
    rate_observations_path = processed / "rate_observations.csv"
    observation_rows = _csv_count(rate_observations_path)
    if observation_rows is not None:
        checks["retained_rate_observation_count_matches_C4"] = observation_rows == c4_rows

    max_workers = min(4, max(1, int(pipeline_cfg.get("max_parallel_rate_files", 1))))
    worker_values = pd.to_numeric(c4.get("worker_pool_size", pd.Series(dtype=str)), errors="coerce").dropna() if not c4.empty else pd.Series(dtype=float)
    actual_workers = int(worker_values.max()) if not worker_values.empty else min(max_workers, max(1, len(c4)))
    self_test = _json(processed / "parallel_processing_selftest.json")

    scorecard = [
        {"track": "1+2", "metric": "Explicit California group PPO plan links classified", "value": selected_plans, "denominator": "", "share": "", "status": "measured", "note": "Plan-link classification does not imply unique plan-level rates."},
        {"track": "1+2", "metric": "Physical rate files selected", "value": c4_complete, "denominator": selected_files, "share": _ratio(c4_complete, selected_files), "status": "measured", "note": f"{int(c2.get('selected_mixed_plan_files', 0) or 0)} selected file(s) are shared with non-PPO links; rate-level PPO attribution remains unresolved."},
        {"track": "1+2", "metric": "Rate observations streamed", "value": c4_rows, "denominator": "", "share": "", "status": "measured", "note": "Large MRF JSON is iterated; output is written incrementally."},
        {"track": "1", "metric": "File quality score / grade", "value": quality_score, "denominator": "100", "share": _ratio(quality_score, 100), "status": quality_grade or "not_scored", "note": "Quality metrics are screening diagnostics, not ground truth."},
        {"track": "2", "metric": "NPI/TIN pairs with an address candidate", "value": candidate_pairs, "denominator": total_pairs, "share": _ratio(candidate_pairs, total_pairs), "status": "candidate_only", "note": "NPPES practice addresses are candidates, not payer-confirmed TIN/address links."},
        {"track": "2", "metric": "Payer-verified NPI/TIN/address pairs", "value": verified_pairs, "denominator": total_pairs, "share": _ratio(verified_pairs, total_pairs), "status": "blocked_missing_payer_directory", "note": "No current payer directory with exact NPI, TIN, address, and effective date was supplied."},
        {"track": "2", "metric": "Workbook rate rows resolved to address candidates", "value": provider_rate_rows, "denominator": provider_rate_rows + unresolved, "share": _ratio(provider_rate_rows, provider_rate_rows + unresolved), "status": "measured_with_unresolved_rows", "note": f"{unresolved} numeric rate observations remain on the unresolved worksheet."},
        {"track": "2", "metric": "Rows with mapped specialty", "value": mapped_rows, "denominator": mapped_rows + unmapped_rows, "share": _ratio(mapped_rows, mapped_rows + unmapped_rows), "status": "measured_with_unmapped_rows", "note": f"{unmapped_rows} rows have no specialty in the supplied crosswalk or NUCC fallback."},
        {"track": "2", "metric": "C4 configured worker pool", "value": actual_workers, "denominator": max_workers, "share": _ratio(actual_workers, max_workers), "status": "bounded_file_parallel" if actual_workers > 1 else "one_selected_file", "note": f"Selected source file count: {len(c4)}; concurrency applies across distinct selected files."},
        {"track": "2", "metric": "Parallel processing self-test", "value": "passed" if self_test.get("status") == "passed" else "not_run_or_failed", "denominator": "2 workers", "share": "", "status": self_test.get("status", "not_run"), "note": "Synthetic two-file fixture; validates concurrent worker overlap and byte-stable output against sequential mode."},
        {"track": "2", "metric": "Answer-key comparison", "value": "not available", "denominator": "answer key not supplied", "share": "", "status": "not_performed", "note": "Internal row-count, uniqueness, and output consistency checks do not substitute for an answer key."},
    ]
    scorecard_path = processed / "track2_scorecard.csv"
    pd.DataFrame(scorecard).to_csv(scorecard_path, index=False)

    config_path = Path(config.get("_config_path", project_root / "config_both_tracks.yaml"))
    if not config_path.is_absolute():
        config_path = project_root / config_path
    source_paths: list[tuple[str, Path]] = []
    source_paths.extend((f"index:{path.name}", path) for path in sorted(dirs["indexes"].glob("*.json")))
    for key in ("mapping_workbook", "nucc_taxonomy_csv", "nppes_csv", "provider_directory_csv"):
        value = sources.get(key)
        if value:
            path = Path(value)
            if not path.is_absolute():
                path = project_root / path
            source_paths.append((key, path))
    for value in sources.get("nppes_additional_csvs", []) or []:
        path = Path(value)
        if not path.is_absolute():
            path = project_root / path
        source_paths.append((f"nppes_additional:{path.name}", path))
    for path in external_sources or []:
        source_paths.append((f"external_source:{path.name}", path))
    for record in c4.to_dict("records"):
        local_path = str(record.get("local_path", "")).strip()
        if local_path:
            path = Path(local_path)
            if not path.is_absolute():
                path = project_root / path
            source_paths.append((f"rate_file:{record.get('file_id', '')}", path))

    source_hashes = []
    seen: set[str] = set()
    for label, path in source_paths:
        key = str(path.resolve()) if path.exists() else str(path)
        if key in seen:
            continue
        seen.add(key)
        source_hashes.append({
            "name": label,
            "basename": path.name,
            "location_scope": "project" if path.exists() and path.resolve().is_relative_to(project_root) else "external_or_missing",
            "project_relative_path": path.resolve().relative_to(project_root).as_posix() if path.exists() and path.resolve().is_relative_to(project_root) else None,
            "available": path.is_file(),
            "bytes": path.stat().st_size if path.is_file() else None,
            "sha256": _hash_file(path) if path.is_file() else None,
            "omitted_reason": "Large source extract is intentionally not bundled; place at configured relative path for a full rerun." if label in {"nppes_csv", "provider_directory_csv"} and not path.is_file() else None,
        })

    code_hashes = []
    for path in sorted((project_root / "pipeline").glob("*.py")):
        if path.name == "c6_build_provisional_rate_file.py":
            continue
        code_hashes.append({"path": path.relative_to(project_root).as_posix(), "sha256": _hash_file(path), "bytes": path.stat().st_size})

    generated_paths = [
        output / "provider_rates.xlsx",
        processed / "c1_manifest.json",
        processed / "c2_manifest.json",
        processed / "c4_file_status.csv",
        processed / "c5_manifest.json",
        processed / "c6_manifest.json",
        processed / "file_quality_metrics.csv",
        processed / "rate_selection_audit.csv",
        processed / "unresolved_rate_observations.csv",
        processed / "provider_address_links.csv",
        scorecard_path,
    ]
    generated_hashes = [
        {"path": path.resolve().relative_to(project_root).as_posix(), "bytes": path.stat().st_size, "sha256": _hash_file(path)}
        for path in generated_paths if path.is_file()
    ]

    validation = {
        "project": "Centene Provider Rate Curation — Tracks 1 and 2",
        "created_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "validation_status": "passed_with_known_requirement_gaps" if all(checks.values()) else "failed",
        "checks": checks,
        "workbook_check": workbook,
        "c4_file_count": len(c4),
        "c4_rows_from_status": c4_rows,
        "rate_observation_csv_rows": observation_rows,
        "parallel_test": self_test,
        "answer_key_validation": "not_performed_answer_key_not_provided",
        "known_gaps": [
            "The selected physical Health Net CA MRF is linked by the index to PPO and non-PPO plans; the supplied sources do not assign each rate to a particular plan.",
            "No current payer provider directory with exact NPI, TIN, practice address, and active/effective dates was supplied; NPPES/CMS addresses remain candidates.",
            f"{unresolved} numeric rate observations have no NPI/TIN address candidate.",
            f"{unmapped_rows} provider-rate rows remain specialty-unmapped.",
            "No answer key was supplied for the Track 2 sample-rate comparison criterion.",
            "The final handoff includes a video outline only per user request; the problem statement's recorded-video deliverable is therefore not included.",
        ],
        "counts": {
            "explicit_ppo_plan_links": selected_plans,
            "selected_physical_rate_files": selected_files,
            "extracted_rate_observations": c4_rows,
            "provider_rate_workbook_rows": provider_rate_rows,
            "unique_workbook_keys": workbook.get("unique_provider_rate_keys"),
            "unresolved_numeric_rate_rows": unresolved,
            "npi_tin_pairs": total_pairs,
            "pairs_with_address_candidates": candidate_pairs,
            "payer_verified_pairs": verified_pairs,
            "unmapped_specialty_rows": unmapped_rows,
        },
    }
    validation_path = processed / "track2_validation.json"
    write_json(validation_path, validation)

    manifest = {
        "project": validation["project"],
        "created_utc": validation["created_utc"],
        "run_date": sources.get("as_of_date"),
        "python_version": platform.python_version(),
        "library_versions": {
            name: _version(name) for name in ("pandas", "ijson", "requests", "openpyxl", "PyYAML")
        },
        "config": {
            "path": config_path.name,
            "sha256": _hash_file(config_path) if config_path.is_file() else None,
        },
        "source_inputs": source_hashes,
        "supplemental_source_evidence": {
            "source": cms.get("source"),
            "source_file_name": cms.get("source_file_name"),
            "dataset_url": cms.get("source_dataset_url"),
            "audit_sha256": cms.get("local_file_sha256"),
            "rows_scanned": cms.get("source_rows_scanned"),
            "note": "This CMS source was reviewed separately as NPI-level context; it is not used as payer TIN/address verification.",
        },
        "pipeline_code": code_hashes,
        "generated_outputs": generated_hashes,
        "source_scope_note": "The large NPPES source extracts and MRF source body are not included in the compact submission archive. Their availability and hashes are recorded where accessible; a fresh full run requires the configured NPPES extracts and network access to the selected MRF.",
        "validation_status": validation["validation_status"],
        "answer_key_comparison": "not_performed_answer_key_not_provided",
    }
    manifest_path = processed / "reproducibility_manifest.json"
    write_json(manifest_path, manifest)
    if logger:
        logger.info("Track 2 audit: %s", validation["validation_status"])
    return {"validation": validation, "manifest": manifest, "scorecard_path": scorecard_path, "validation_path": validation_path, "manifest_path": manifest_path}


def _version(name: str) -> str | None:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return None


def main() -> None:
    parser = argparse.ArgumentParser(description="Create Track 2 validation and reproducibility evidence")
    parser.add_argument("--config", default="config_both_tracks.yaml")
    parser.add_argument("--external-source", action="append", default=[], help="Additional source file to hash; can be repeated. Absolute paths are never written into the manifest.")
    args = parser.parse_args()
    run(load_config(args.config), external_sources=[Path(value).resolve() for value in args.external_source])


if __name__ == "__main__":
    main()
